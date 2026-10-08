import openpyxl
import pandas as pd

from datos_excel import (
  COLUMNAS_PERFIL, COLUMNAS_WISHLIST,
  cargar_datos, guardar_excel, normalizar_perfil, normalizar_wishlist,
  excel_en_bytes, filtrar_tabla, preparar_edicion, resumir_cambios, validar_wishlist,
)


def _wishlist_antigua():
  """Hoja como estaba antes de las columnas nuevas, con los defectos reales de tu Excel."""
  return pd.DataFrame({
    "Articulo": ["Medias ", "\n Mando", "Power bank"],
    "Tipo": ["Deseo ", "Deseo", None],
    "Costo Estimado": [50.0, 2.0, None],
    "Estado": ["Comprado", None, None],
    "Notas": [None, None, None],
  })


def test_normalizar_wishlist_quita_espacios_y_saltos_de_linea():
  limpios = normalizar_wishlist(_wishlist_antigua())

  assert limpios["Articulo"].tolist() == ["Medias", "Mando", "Power bank"]
  assert limpios["Tipo"].iloc[0] == "Deseo"


def test_normalizar_wishlist_rellena_vacios_sin_inventar_costos():
  limpios = normalizar_wishlist(_wishlist_antigua())

  assert limpios["Tipo"].iloc[2] == "Sin clasificar"
  assert limpios["Estado"].tolist() == ["Comprado", "Pendiente", "Pendiente"]
  assert limpios["Costo Estimado"].isna().iloc[2]


def test_normalizar_wishlist_agrega_columnas_nuevas_con_defectos():
  limpios = normalizar_wishlist(_wishlist_antigua())

  assert list(limpios.columns[: len(COLUMNAS_WISHLIST)]) == COLUMNAS_WISHLIST
  assert (limpios["Urgencia"] == 3).all()
  assert (limpios["Frecuencia"] == "Único").all()


def test_normalizar_wishlist_no_modifica_el_original():
  original = _wishlist_antigua()
  normalizar_wishlist(original)

  assert original["Articulo"].iloc[0] == "Medias "


def test_normalizar_wishlist_acepta_dataframe_vacio():
  limpios = normalizar_wishlist(pd.DataFrame(columns=COLUMNAS_WISHLIST))

  assert limpios.empty
  assert list(limpios.columns) == COLUMNAS_WISHLIST


def test_normalizar_perfil_limpia_cabeceras():
  sucio = pd.DataFrame(columns=["Ingreso Neto Mensual ", "Gastos fijos ", "Deseos ya gastados ",
                                "Ahorro actual", "Meses de Fondo de emergencia objetivo"])

  assert list(normalizar_perfil(sucio).columns) == COLUMNAS_PERFIL


def _lista_normalizada():
  return normalizar_wishlist(pd.DataFrame({
    "Articulo": ["Medias", "Mouse", "Pedalera"],
    "Tipo": ["Necesidad", "Necesidad", "Deseo"],
    "Costo Estimado": [30.0, 60.0, None],
    "Estado": ["Pendiente", "Pendiente", "Pendiente"],
  }))


def test_validar_wishlist_acepta_una_lista_correcta():
  assert validar_wishlist(_lista_normalizada()) == []


def test_validar_wishlist_detecta_nombre_vacio_repetido_y_costo_cero():
  lista = _lista_normalizada()
  lista.loc[0, "Articulo"] = "mouse"        # repetido con "Mouse" aunque cambie la mayuscula
  lista.loc[2, "Articulo"] = ""            # sin nombre
  lista.loc[1, "Costo Estimado"] = 0.0     # costo cero

  errores = validar_wishlist(lista)

  assert any("Faltan nombres" in e and "fila 3" in e for e in errores)
  assert any("repetidos" in e for e in errores)
  assert any("mayor que cero" in e for e in errores)


def test_validar_wishlist_detecta_puntajes_fuera_de_rango():
  lista = _lista_normalizada()
  lista.loc[0, "Urgencia"] = 9

  errores = validar_wishlist(lista)

  assert any("Urgencia" in e and "Medias" in e for e in errores)


def test_resumir_cambios_cuenta_agregadas_borradas_y_modificadas():
  original = _lista_normalizada()
  editado = original.drop(index=1)                       # borra Mouse
  editado.loc[0, "Costo Estimado"] = 35.0                # modifica Medias
  editado.loc[7] = editado.loc[0]                        # agrega una fila (indice nuevo)

  assert resumir_cambios(original, editado) == {"agregadas": 1, "borradas": 1, "modificadas": 1}


def test_resumir_cambios_sin_cambios_y_vacio_contra_nan():
  original = _lista_normalizada()
  editado = original.copy()
  editado["Notas"] = None   # NaN contra None no es un cambio

  assert resumir_cambios(original, editado) == {"agregadas": 0, "borradas": 0, "modificadas": 0}


def test_preparar_edicion_da_fecha_de_hoy_a_las_filas_nuevas_y_defectos():
  original = _lista_normalizada()
  hoy = pd.Timestamp("2026-10-08")
  editado = original.copy()
  editado.loc[9] = [" Teclado ", None, 90.0, None, None, None, None, None, None, None][: len(editado.columns)]

  limpia, errores, cambios = preparar_edicion(original, editado, hoy)

  nueva = limpia[limpia["Articulo"] == "Teclado"].iloc[0]
  assert errores == []
  assert cambios["agregadas"] == 1
  assert nueva["Fecha de alta"] == hoy
  assert nueva["Tipo"] == "Sin clasificar" and nueva["Estado"] == "Pendiente" and nueva["Urgencia"] == 3
  assert list(limpia.index) == [0, 1, 2, 3]


def test_preparar_edicion_convierte_costos_en_cero_a_por_cotizar_y_avisa():
  original = _lista_normalizada()
  original.loc[0, "Costo Estimado"] = 0.0     # dato viejo con costo 0
  editado = original.copy()
  editado.loc[1, "Costo Estimado"] = 65.0     # el usuario solo cambia otro precio

  limpia, errores, cambios = preparar_edicion(original, editado)

  assert errores == []
  assert pd.isna(limpia.loc[0, "Costo Estimado"])
  assert cambios["costos_en_cero"] == 1 and cambios["modificadas"] == 1


def test_preparar_edicion_sin_cambios_no_reporta_cambios_reales():
  original = _lista_normalizada()

  _, _, cambios = preparar_edicion(original, original.copy())

  assert (cambios["agregadas"], cambios["borradas"], cambios["modificadas"]) == (0, 0, 0)


def test_filtrar_tabla_combina_filtros_y_busqueda_sin_modificar_el_original():
  lista = _lista_normalizada()

  solo_necesidades = filtrar_tabla(lista, {"Tipo": ["Necesidad"]})
  buscar_mou = filtrar_tabla(lista, texto=" MOU ")
  combinado = filtrar_tabla(lista, {"Tipo": ["Necesidad"], "Estado": ["Comprado"]})

  assert solo_necesidades["Articulo"].tolist() == ["Medias", "Mouse"]
  assert buscar_mou["Articulo"].tolist() == ["Mouse"]
  assert combinado.empty
  assert len(lista) == 3


def test_filtrar_tabla_con_listas_vacias_no_filtra():
  lista = _lista_normalizada()

  assert len(filtrar_tabla(lista, {"Tipo": [], "Estado": []}, "")) == 3
  assert len(filtrar_tabla(lista)) == 3


def test_excel_en_bytes_genera_un_libro_valido_con_las_tablas():
  from io import BytesIO
  import openpyxl

  contenido = excel_en_bytes({"WishList": _lista_normalizada(), "Perfil": pd.DataFrame(columns=COLUMNAS_PERFIL)})

  libro = openpyxl.load_workbook(BytesIO(contenido))
  assert libro.sheetnames == ["WishList", "Perfil"]
  assert dict(libro["WishList"].tables.items()) == {"TablaWishList": "A1:I4"}
  assert dict(libro["Perfil"].tables.items()) == {}   # hoja sin filas: queda como rango
  assert pd.read_excel(BytesIO(contenido))["Articulo"].tolist() == ["Medias", "Mouse", "Pedalera"]


def test_guardar_y_cargar_conserva_ambas_hojas_con_una_tabla_cada_una(tmp_path):
  ruta = tmp_path / "prueba.xlsx"
  perfil = pd.DataFrame([[1000.0, 400.0, 100.0, 500.0, 3.0]], columns=COLUMNAS_PERFIL)
  hojas = {"WishList": normalizar_wishlist(_wishlist_antigua()), "Perfil": perfil}

  guardar_excel(hojas, ruta)
  guardar_excel(cargar_datos(ruta), ruta)  # segundo guardado: no debe duplicar tablas

  libro = openpyxl.load_workbook(ruta)
  assert libro.sheetnames == ["WishList", "Perfil"]
  assert dict(libro["WishList"].tables.items()) == {"TablaWishList": f"A1:I4"}
  assert dict(libro["Perfil"].tables.items()) == {"TablaPerfil": "A1:E2"}
  assert cargar_datos(ruta)["Perfil"]["Ingreso Neto Mensual"].iloc[0] == 1000.0


def test_hoja_sin_filas_queda_sin_tabla(tmp_path):
  ruta = tmp_path / "vacio.xlsx"

  guardar_excel(cargar_datos(ruta), ruta)

  libro = openpyxl.load_workbook(ruta)
  assert dict(libro["Perfil"].tables.items()) == {}
