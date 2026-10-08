#importar librerias necesarias para el sistema
from pathlib import Path
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.table import TableStyleInfo
import streamlit as st


#ruta del archivo excel junto a este script (no depende de la carpeta desde donde se lance streamlit)
archivo = Path(__file__).parent / "TablaWishList.xlsx"

# Configuración de la página
st.set_page_config(
    page_title="WishList", page_icon="📊", layout="centered"
)

st.title("📝 WishList")
st.markdown("Sistema para organizar tus inversiones y prioridades")

COLUMNAS_WISHLIST = ["Articulo", "Tipo", "Costo Estimado", "Estado", "Notas"]
COLUMNAS_PERFIL = [
    "Ingreso Mensual", "Gastos Fijos", "Deseos Gastados Mes",
    "Ahorro Actual", "Meses Fondo Objetivo",
]

#cargar todas las hojas del excel en un diccionario {nombre_de_hoja: DataFrame}
def cargar_datos():
  vacias = {
      "WishList": pd.DataFrame(columns=COLUMNAS_WISHLIST),
      "Perfil": pd.DataFrame(columns=COLUMNAS_PERFIL),
  }
  if not archivo.exists():
    return vacias
  hojas = pd.read_excel(archivo, sheet_name=None)
  # las hojas leidas pisan a las vacias; si falta alguna se usa la vacia
  return {**vacias, **hojas}


#guardar el DataFrame en Excel dejando los datos como Tabla oficial (no como rango)
def agregar_tabla(hoja, datos, nombre):
  # Excel no acepta una tabla sin filas de datos: se deja la hoja como rango
  if datos.empty:
    return

  # Rango exacto: de A1 hasta la ultima columna y la ultima fila (cabecera + datos)
  ultima_columna = get_column_letter(datos.shape[1])
  rango = f"A1:{ultima_columna}{len(datos) + 1}"

  tabla = Table(displayName=nombre, ref=rango)
  tabla.tableStyleInfo = TableStyleInfo(
      name="TableStyleMedium9",
      showFirstColumn=False,
      showLastColumn=False,
      showRowStripes=True,
  )
  hoja.add_table(tabla)


def guardar_excel(hojas, ruta):
  # to_excel crea el libro desde cero: no queda ninguna tabla anterior, asi que no puede duplicarse
  with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
    for nombre_hoja, datos in hojas.items():
      datos.to_excel(writer, sheet_name=nombre_hoja, index=False)
      # nombre unico por hoja y sin espacios: TablaWishList, TablaPerfil
      nombre_tabla = f"Tabla{nombre_hoja.replace(' ', '')}"
      agregar_tabla(writer.sheets[nombre_hoja], datos, nombre_tabla)

hojas = cargar_datos()
df = hojas["WishList"]

# Formulario de entrada para los campos
with st.form("entry_form", clear_on_submit=True):

  st.subheader("Nuevo Articulo")
  col1, col2 = st.columns(2)

  with col1:

  #Insercion de campos de entrada para el formulario de datos

    # Campo de entrada para el nombre del artículo
    articulo = st.text_input("Articulo")

    # Campo de selección para el tipo de artículo
    tipo = st.selectbox( "Tipo", ["Deseo", "Necesidad"] )
  with col2:

    # Campo de entrada para el costo estimado
    costo_estimado = st.number_input("Costo Estimado", value=0.0, min_value=0.)

    # Campo de selección para el estado del artículo
    estado = st.selectbox("Estado", ["Pendiente", "Comprado"])

    # Campo de entrada para las notas del artículo
    notas = st.text_area("Notas")

   #Boton para guardar los datos en el archivo excel
  submit_button = st.form_submit_button(
      label="Guardar", use_container_width=True
  )

  if submit_button:

    # Reunir todos los problemas antes de decidir si se guarda
    errores = []
    if articulo.strip() == "":
      errores.append("Por favor, completa al menos el campo de Articulo.")
    if costo_estimado <= 0:
      errores.append("El costo estimado debe ser mayor que cero.")

    if errores:
      for error in errores:
        st.error(error)
    else:
      # Nuevo registro
      nuevo_articulo = pd.DataFrame([{
          "Articulo": articulo.strip(),
          "Tipo": tipo,
          "Costo Estimado": costo_estimado,
          "Estado": estado,
          "Notas": notas,
      }])

      # Concatenar en una copia: df solo cambia si el guardado sale bien
      df_actualizado = pd.concat([df, nuevo_articulo], ignore_index=True)

      try:
        # se reenvian TODAS las hojas: lo que no se escribe se pierde
        guardar_excel({**hojas, "WishList": df_actualizado}, archivo)
      except PermissionError:
        st.error("El Excel está abierto o OneDrive lo bloquea. Ciérralo y vuelve a guardar.")
      else:
        df = df_actualizado
        st.success("¡Datos guardados exitosamente en el Excel!")

# Visualización rápida de los datos actuales
st.divider()
st.subheader("Vista previa de los datos actuales")
st.dataframe(df, use_container_width=True)

# Botón de descarga directa
if not df.empty:
  with open(archivo, "rb") as f:
    st.download_button(
        label="📥 Descargar Excel Actualizado",
        data=f,
        file_name=archivo.name,
        mime=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
