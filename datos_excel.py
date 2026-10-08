"""Lectura, limpieza y escritura del Excel central (hojas WishList y Perfil).

Este modulo no usa Streamlit: asi se puede probar con pytest sin abrir la app.
"""
from io import BytesIO

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COLUMNAS_WISHLIST = [
  "Articulo", "Tipo", "Costo Estimado", "Estado", "Notas",
  "Frecuencia", "Urgencia", "Valor", "Fecha de alta",
]
COLUMNAS_PERFIL = [
  "Ingreso Neto Mensual", "Gastos fijos", "Deseos ya gastados",
  "Ahorro actual", "Meses de Fondo de emergencia objetivo",
]

# Todos los montos estan en dolares (USD). Los nombres de columna no cambian: solo se muestra el simbolo.
MONEDA = "USD"
COLUMNAS_DINERO = {
  "Costo Estimado", "Compromiso Mensual",
  "Ingreso Neto Mensual", "Gastos fijos", "Deseos ya gastados", "Ahorro actual",
}
FORMATO_DINERO_EXCEL = '"$"#,##0.00'

TIPOS = ["Deseo", "Necesidad"]
TIPOS_VALIDOS = TIPOS + ["Sin clasificar"]  # "Sin clasificar" es el valor por defecto de los datos viejos
ESTADOS = ["Pendiente", "Comprado"]
FRECUENCIAS = ["Único", "Mensual", "Anual"]

COLUMNAS_TEXTO = ["Articulo", "Tipo", "Estado", "Frecuencia"]
COLUMNAS_PUNTAJE = ["Urgencia", "Valor"]
PUNTAJE_POR_DEFECTO = 3
TEXTO_POR_DEFECTO = {
  "Tipo": "Sin clasificar",
  "Estado": "Pendiente",
  "Frecuencia": "Único",
}


def _con_columnas(datos, columnas):
  """Devuelve una copia con las columnas esperadas primero; agrega las que falten (vacias)
  y conserva las extra que el usuario haya puesto en Excel."""
  extras = [c for c in datos.columns if c not in columnas]
  return datos.reindex(columns=columnas + extras)


def normalizar_wishlist(datos):
  """Limpia la hoja WishList: columnas completas, textos sin espacios y valores por defecto."""
  limpios = _con_columnas(datos, COLUMNAS_WISHLIST)

  for columna in COLUMNAS_TEXTO:
    # astype("string") evita el error de .str cuando la columna viene vacia; los vacios siguen vacios
    limpios[columna] = limpios[columna].astype("string").str.strip()
  for columna, defecto in TEXTO_POR_DEFECTO.items():
    limpios[columna] = limpios[columna].fillna(defecto)

  for columna in COLUMNAS_PUNTAJE:
    # errors="coerce" convierte lo que no sea numero en vacio, y luego se rellena
    numeros = pd.to_numeric(limpios[columna], errors="coerce")
    limpios[columna] = numeros.fillna(PUNTAJE_POR_DEFECTO).astype(int)

  limpios["Fecha de alta"] = pd.to_datetime(limpios["Fecha de alta"], errors="coerce")
  return limpios


def _nombres_o_fila(datos):
  """Etiqueta de cada fila para mensajes de error: el nombre, o 'fila N' si no tiene."""
  nombres = datos["Articulo"].astype(object)
  return [
    n if isinstance(n, str) and n.strip() else f"fila {posicion + 1}"
    for posicion, n in enumerate(nombres)
  ]


def validar_wishlist(datos):
  """Revisa una WishList ya normalizada y devuelve la lista de problemas (vacia si esta bien)."""
  errores = []
  etiquetas = pd.Series(_nombres_o_fila(datos), index=datos.index)
  nombres = datos["Articulo"]

  sin_nombre = nombres.isna() | nombres.eq("").fillna(False)
  if sin_nombre.any():
    errores.append(f"Faltan nombres en: {', '.join(etiquetas[sin_nombre])}. Escribe el nombre o borra la fila.")

  con_nombre = nombres[~sin_nombre]
  repetidos = con_nombre[con_nombre.str.lower().duplicated(keep=False)]
  if not repetidos.empty:
    errores.append(f"Artículos repetidos: {', '.join(sorted(set(repetidos)))}.")

  costos = pd.to_numeric(datos["Costo Estimado"], errors="coerce")
  if (costos <= 0).any():
    errores.append(
      f"El costo debe ser mayor que cero (o vacío si falta cotizar): {', '.join(etiquetas[costos <= 0])}."
    )

  for columna in COLUMNAS_PUNTAJE:
    fuera = ~datos[columna].between(1, 5)
    if fuera.any():
      errores.append(f"{columna} debe estar entre 1 y 5: {', '.join(etiquetas[fuera])}.")

  for columna, validos in (("Tipo", TIPOS_VALIDOS), ("Estado", ESTADOS), ("Frecuencia", FRECUENCIAS)):
    invalidos = ~datos[columna].isin(validos)
    if invalidos.any():
      errores.append(f"{columna} no válido en: {', '.join(etiquetas[invalidos])}.")
  return errores


def _columna_cambio(antes, despues):
  """True en las filas donde el valor cambio. NaN contra vacio no cuenta como cambio."""
  if pd.api.types.is_numeric_dtype(antes) and pd.api.types.is_numeric_dtype(despues):
    iguales = (antes == despues) | (antes.isna() & despues.isna())
    return ~iguales
  def como_texto(serie):
    return serie.astype(object).where(serie.notna(), "").astype(str).str.strip()

  return como_texto(antes) != como_texto(despues)


def resumir_cambios(original, editado):
  """Cuenta filas agregadas, borradas y modificadas entre la lista original y la editada."""
  mantenidas = editado.index.intersection(original.index)
  # la fecha de alta no se edita a mano: se ignora para no marcar cambios falsos
  columnas = [c for c in original.columns if c in editado.columns and c != "Fecha de alta"]

  modificadas = pd.Series(False, index=mantenidas)
  for columna in columnas:
    modificadas |= _columna_cambio(original.loc[mantenidas, columna], editado.loc[mantenidas, columna])
  return {
    "agregadas": len(editado) - len(mantenidas),
    "borradas": len(original) - len(mantenidas),
    "modificadas": int(modificadas.sum()),
  }


def preparar_edicion(original, editado, hoy=None):
  """Convierte la tabla editada en una WishList lista para guardar.

  Devuelve (lista_limpia, errores, cambios). Las filas nuevas reciben la fecha de alta de hoy.
  """
  hoy = hoy if hoy is not None else pd.Timestamp.today().normalize()
  cambios = resumir_cambios(original, editado)

  trabajo = editado.copy()
  es_nueva = ~trabajo.index.isin(original.index)
  if es_nueva.any():
    trabajo["Fecha de alta"] = trabajo["Fecha de alta"].astype("datetime64[ns]")
    trabajo.loc[es_nueva, "Fecha de alta"] = hoy
  limpia = normalizar_wishlist(trabajo).reset_index(drop=True)

  # un costo 0 significa "aun no cotizado" (asi lo trata el semaforo): se guarda vacio, y se avisa.
  # Sin esto, los 0 de datos viejos bloquearian cualquier edicion no relacionada.
  costos = pd.to_numeric(limpia["Costo Estimado"], errors="coerce")
  en_cero = costos == 0
  limpia["Costo Estimado"] = costos.mask(en_cero)
  cambios["costos_en_cero"] = int(en_cero.sum())
  return limpia, validar_wishlist(limpia), cambios


def normalizar_perfil(datos):
  """Quita espacios de los nombres de columna (una cabecera con espacio final es una trampa en Excel y DAX)."""
  limpios = datos.rename(columns=lambda nombre: str(nombre).strip())
  return _con_columnas(limpios, COLUMNAS_PERFIL)


def cargar_datos(ruta):
  """Carga todas las hojas en un diccionario {nombre_de_hoja: DataFrame}."""
  vacias = {
    "WishList": pd.DataFrame(columns=COLUMNAS_WISHLIST),
    "Perfil": pd.DataFrame(columns=COLUMNAS_PERFIL),
  }
  # las hojas leidas pisan a las vacias; si falta alguna se usa la vacia
  hojas = pd.read_excel(ruta, sheet_name=None) if ruta.exists() else {}
  todas = {**vacias, **hojas}
  return {
    **todas,
    "WishList": normalizar_wishlist(todas["WishList"]),
    "Perfil": normalizar_perfil(todas["Perfil"]),
  }


def agregar_tabla(hoja, datos, nombre):
  """Convierte el rango de datos de la hoja en una Tabla oficial de Excel."""
  # Excel no acepta una tabla sin filas de datos: se deja la hoja como rango
  if datos.empty:
    return

  # Rango exacto: de A1 hasta la ultima columna y la ultima fila (cabecera + datos)
  ultima_columna = get_column_letter(datos.shape[1])
  rango = f"A1:{ultima_columna}{len(datos) + 1}"

  tabla = Table(displayName=nombre, ref=rango)
  tabla.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium9",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
  )
  hoja.add_table(tabla)


def aplicar_formato_dinero(hoja):
  """Muestra en dolares (USD) las columnas de dinero de la hoja. Solo cambia el formato, no el valor."""
  for celda_cabecera in hoja[1]:
    if celda_cabecera.value in COLUMNAS_DINERO:
      for fila in hoja.iter_rows(min_row=2, min_col=celda_cabecera.column, max_col=celda_cabecera.column):
        fila[0].number_format = FORMATO_DINERO_EXCEL


def filtrar_tabla(datos, filtros=None, texto="", columna_texto="Articulo"):
  """Devuelve las filas que cumplen los filtros, sin modificar la tabla original.

  filtros: {columna: [valores permitidos]}. Una lista vacia no filtra esa columna.
  texto: busca (sin importar mayusculas) dentro de columna_texto.
  """
  mascara = pd.Series(True, index=datos.index)
  for columna, valores in (filtros or {}).items():
    if valores:
      mascara &= datos[columna].isin(valores)
  if texto.strip():
    contiene = datos[columna_texto].astype("string").str.contains(texto.strip(), case=False, regex=False, na=False)
    mascara &= contiene
  return datos[mascara]


def excel_en_bytes(hojas):
  """Genera el Excel completo en memoria (para el boton de descarga), sin leer ni escribir el disco."""
  buffer = BytesIO()
  guardar_excel(hojas, buffer)
  return buffer.getvalue()


def guardar_excel(hojas, ruta):
  """Escribe todas las hojas, cada una como Tabla. Crea el libro desde cero:
  no queda ninguna tabla anterior, asi que no puede duplicarse.
  `ruta` puede ser una ruta de archivo o un BytesIO."""
  with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
    for nombre_hoja, datos in hojas.items():
      datos.to_excel(writer, sheet_name=nombre_hoja, index=False)
      hoja = writer.sheets[nombre_hoja]
      aplicar_formato_dinero(hoja)
      # nombre unico por hoja y sin espacios: TablaWishList, TablaPerfil
      nombre_tabla = f"Tabla{nombre_hoja.replace(' ', '')}"
      agregar_tabla(hoja, datos, nombre_tabla)
